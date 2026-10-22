import io
import json
from datetime import datetime, timedelta, timezone
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
import main
import models
import extended_models as ex
from database import Base, get_db
from auth import hash_password
from services.intelligence import account_risk, level, network_scan, shop_trust


@pytest.fixture
def env():
    engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory=sessionmaker(bind=engine)
    def override():
        with factory() as db:yield db
    main.app.dependency_overrides[get_db]=override
    with factory() as db:
        for role in ('buyer','seller','admin','engineer'):
            db.add(models.User(email=f'{role}@example.com',password_hash=hash_password('Password123'),role=role))
        db.commit()
    with TestClient(main.app) as client:
        headers={}
        for role in ('buyer','seller','admin','engineer'):
            r=client.post('/api/auth/login',json={'email':f'{role}@example.com','password':'Password123'})
            assert r.status_code==200,r.text
            headers[role]={'Authorization':f"Bearer {r.json()['access_token']}"}
        yield client,headers,factory
    main.app.dependency_overrides.clear();engine.dispose()


def create_review(env,content='I used this cable for three months and it broke last week.'):
    c,h,_=env
    p=c.post('/api/products',headers=h['seller'],json={'name':'Cable','price':10,'category':'Electronics'}).json()
    denied=c.post('/api/reviews',headers=h['buyer'],json={'product_id':p['id'],'content':content,'rating':3})
    assert denied.status_code==403
    order=c.post('/api/workspace/purchases',headers=h['buyer'],json={'product_id':p['id']}).json()
    assert c.patch(f"/api/workspace/purchases/{order['id']}",headers=h['seller'],json={'status':'completed'}).status_code==200
    response=c.post('/api/reviews',headers=h['buyer'],json={'product_id':p['id'],'content':content,'rating':3})
    assert response.status_code==201,response.text
    return response.json(),p


@pytest.mark.parametrize('score,expected',[(0,'low'),(50,'low'),(50.01,'moderate'),(69.99,'moderate'),(70,'high'),(100,'high')])
def test_threshold_boundaries(score,expected):
    assert level(score,{'normal_max':50,'suspicious_min':70})==expected


def test_real_model_lime_history_and_latest_statistics(env):
    c,h,dbf=env
    review,product=create_review(env)
    p=review['prediction'];assert p['model_name']=='tfidf_svm_calibrated'
    assert 0<p['fake_probability']<1
    assert p['fake_probability']+p['real_probability']==pytest.approx(1)
    assert c.put(f"/api/reviews/{review['id']}",headers=h['buyer'],json={'content':'The cable is amazing and really wonderful.'}).status_code==200
    with dbf() as db:assert db.query(models.Prediction).filter_by(review_id=review['id']).count()==2
    stats=c.get('/api/statistics/overview',headers=h['admin']).json()
    assert stats['total_reviews']==1 and stats['fake_reviews']+stats['real_reviews']==1
    assert c.post('/api/workspace/analyze',headers=h['buyer'],json={'content':'After two weeks the cable stopped charging my phone.'}).status_code==200
    assert len(c.get('/api/workspace/analyses',headers=h['buyer']).json())==1
    explain=c.post('/api/workspace/explain',headers=h['buyer'],json={'content':'After two weeks the cable stopped charging my phone.'})
    assert explain.status_code==200,explain.text
    assert explain.json()['method']=='LIME' and explain.json()['features']


def test_report_resolution_permissions_and_manual_override(env):
    c,h,_=env;review,_=create_review(env)
    endpoint=f"/api/reviews/{review['id']}/report"
    r=c.post(endpoint,headers=h['buyer'],json={'reason':'spam','description':'Repeated text'})
    assert r.status_code==201,r.text
    assert c.post(endpoint,headers=h['buyer'],json={'reason':'spam'}).status_code==409
    rid=r.json()['id'];url=f'/api/workspace/reports/{rid}'
    body={'status':'resolved','note':'Confirmed and reviewed by moderator'}
    assert c.patch(url,headers=h['buyer'],json=body).status_code==403
    assert c.patch(url,headers=h['admin'],json=body).status_code==200
    assert c.get('/api/workspace/reports',headers=h['buyer']).json()[0]['resolution']['note']==body['note']
    assert c.post(f"/api/workspace/reviews/{review['id']}/moderate",headers=h['admin'],json={'label':'real','reason':'Reviewed original experience'}).status_code==200
    audit=c.get(f"/api/workspace/reviews/{review['id']}/audit",headers=h['admin']).json()
    assert len(audit['predictions'])==1 and audit['moderations'][0]['label']=='real'
    assert c.get('/api/workspace/network',headers=h['buyer']).status_code==403


def test_short_reviews_and_account_ratio(env):
    c,h,dbf=env
    for content in ['a','b','c','d','e','f','g','A cable with a broken pin after two months','The headphones arrived in a cardboard box','Two adapters included and both work correctly']:
        create_review(env,content)
    with dbf() as db:
        # Move reviews to different days so the test isolates short-content signals.
        rows=db.query(models.Review).all()
        for i,r in enumerate(rows):r.created_at=datetime(2026,1,1)+timedelta(days=i)
        db.commit();uid=rows[0].user_id
        risk=account_risk(db,uid)
        assert risk['score']==70 and risk['level']=='high'
        assert risk['abnormal_reviews']==7


def test_network_requires_multiple_shared_products_and_shop_trust(env):
    c,h,factory=env
    with factory() as db:
        seller=db.query(models.User).filter_by(role='seller').first();buyer=db.query(models.User).filter_by(role='buyer').first()
        other=models.User(email='other@example.com',password_hash=hash_password('Password123'),role='buyer');db.add(other);db.flush()
        for i in range(2):
            p=models.Product(product_code=f'NETWORK-{i}',name=f'Product {i}',price=1,shop_id=seller.id);db.add(p);db.flush()
            for u in (buyer,other):db.add(models.Review(product_id=p.id,user_id=u.id,rating=5,content='same promotional statement',created_at=datetime(2026,1,1)))
        db.commit()
        graph=network_scan(db)
        assert len(graph['clusters'])==1 and len(graph['clusters'][0]['members'])==2
        trust=shop_trust(db,seller.id,True)
        assert trust['trust_score'] is None and trust['pending_reviews']==4


def test_password_revokes_sessions_and_reset_expiry(env):
    c,h,_=env
    assert c.post('/api/workspace/password',headers=h['buyer'],json={'current_password':'wrong','new_password':'ChangedPassword123'}).status_code==400
    assert c.post('/api/workspace/password',headers=h['buyer'],json={'current_password':'Password123','new_password':'ChangedPassword123'}).status_code==200
    assert c.get('/api/auth/me',headers=h['buyer']).status_code==401
    assert c.post('/api/auth/login',json={'email':'buyer@example.com','password':'ChangedPassword123'}).status_code==200
    assert c.post('/api/workspace/forgot-password',json={'email':'nonexistent@example.com'}).status_code==200
    assert c.post('/api/workspace/reset-password',json={'token':'x'*30,'new_password':'PasswordNew123'}).status_code==400


def test_export_scope_and_spreadsheet_formula_injection(env):
    c,h,_=env;create_review(env,'=SUM(1,2)')
    response=c.get('/api/workspace/exports/xlsx',headers=h['seller'])
    assert response.status_code==200
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(response.content))
    assert wb.active['F2'].data_type=='s' and wb.active['F2'].value=='=SUM(1,2)'
    assert c.get('/api/workspace/exports/xlsx',headers=h['buyer']).status_code==403
    response=c.get('/api/workspace/exports/pdf',headers=h['seller'])
    assert response.status_code==200 and response.content.startswith(b'%PDF')


def test_review_update_rolls_back_if_analysis_followup_fails(env,monkeypatch):
    c,h,factory=env;review,_=create_review(env)
    import services.intelligence as intelligence
    def fail(*args,**kwargs):raise RuntimeError('simulated storage failure')
    monkeypatch.setattr(intelligence,'account_risk',fail)
    with pytest.raises(RuntimeError):
        c.put(f"/api/reviews/{review['id']}",headers=h['buyer'],json={'content':'New text must not be partially committed'})
    with factory() as db:
        assert db.get(models.Review,review['id']).content==review['content']
        assert db.query(models.Prediction).filter_by(review_id=review['id']).count()==1


def test_stock_failure_keeps_order_pending_and_preserves_history(env):
    c,h,factory=env
    product=c.post('/api/products',headers=h['seller'],json={'name':'Last cable','price':10}).json()
    pid=product['id']
    assert c.put(f'/api/workspace/products/{pid}/extra',headers=h['seller'],json={'stock':1}).status_code==200
    orders=[c.post('/api/workspace/purchases',headers=h['buyer'],json={'product_id':pid}).json() for _ in range(2)]
    assert c.patch(f"/api/workspace/purchases/{orders[0]['id']}",headers=h['seller'],json={'status':'completed'}).status_code==200
    assert c.patch(f"/api/workspace/purchases/{orders[1]['id']}",headers=h['seller'],json={'status':'completed'}).status_code==409
    with factory() as db:
        assert db.get(ex.ProductExtra,pid).stock==0
        assert db.get(ex.Purchase,orders[1]['id']).status=='pending'
    assert c.delete(f'/api/products/{pid}',headers=h['seller']).status_code==409


def test_reset_token_can_only_be_used_once(env):
    c,h,factory=env
    assert c.post('/api/workspace/forgot-password',json={'email':'buyer@example.com'}).status_code==200
    with factory() as db:
        token=json.loads(db.query(ex.Outbox).one().payload)['token']
    payload={'token':token,'new_password':'ResetPassword123'}
    assert c.post('/api/workspace/reset-password',json=payload).status_code==200
    assert c.post('/api/workspace/reset-password',json=payload).status_code==400
    assert c.get('/api/auth/me',headers=h['buyer']).status_code==401


def test_report_datasets_and_role_scope(env):
    from openpyxl import load_workbook
    c,h,_=env
    review,product=create_review(env)
    for kind in ('reviews','shops','categories','clusters','models'):
        response=c.get(f'/api/workspace/exports/xlsx?kind={kind}',headers=h['admin'])
        assert response.status_code==200,response.text
        sheet=load_workbook(io.BytesIO(response.content)).active
        assert sheet.max_column >= 4
        pdf=c.get(f'/api/workspace/exports/pdf?kind={kind}',headers=h['admin'])
        assert pdf.status_code==200 and pdf.content.startswith(b'%PDF')
    assert c.get('/api/workspace/exports/xlsx?kind=clusters',headers=h['seller']).status_code==403
    assert c.get('/api/workspace/exports/xlsx?kind=models',headers=h['buyer']).status_code==403
    assert c.post(f"/api/workspace/reviews/{review['id']}/moderate",headers=h['admin'],json={'label':'fake','reason':'Verified manual test label'}).status_code==200
    filtered=c.get(f"/api/reviews/product/{product['id']}?label=fake").json()
    assert filtered['total']==1 and filtered['items'][0]['effective_label']=='fake'


def test_missing_model_does_not_silently_return_heuristic(env,monkeypatch):
    from ml.predictor import predictor
    c,h,_=env
    monkeypatch.setattr(predictor,'model',None)
    r=c.post('/api/workspace/analyze',headers=h['buyer'],json={'content':'A wonderful product'})
    assert r.status_code==503


def test_ab_assignments_are_sticky_and_only_verified_labels_are_scored(env,monkeypatch):
    import ml.predictor as model_module
    from fastapi import HTTPException
    from types import SimpleNamespace
    def selected(filename):
        if filename not in ('a.joblib','b.joblib'):raise HTTPException(400,'Unknown model')
        fake=.9 if filename=='b.joblib' else .1
        return SimpleNamespace(predict=lambda text:dict(predicted_label='fake' if fake>.5 else 'real',fake_probability=fake,real_probability=1-fake,confidence_score=.9,explanation='Controlled test model',model_name=filename,model_version='test'))
    monkeypatch.setattr(model_module,'version_predictor',selected)
    c,h,_=env
    assert c.post('/api/workspace/models/experiments',headers=h['buyer'],json={'model_a':'a.joblib','model_b':'b.joblib','traffic_b':20}).status_code==403
    response=c.post('/api/workspace/models/experiments',headers=h['admin'],json={'model_a':'a.joblib','model_b':'b.joblib','traffic_b':20})
    assert response.status_code==200,response.text
    first,_=create_review(env);second,_=create_review(env)
    assert first['prediction']['model_name']==second['prediction']['model_name']
    summary=c.get('/api/workspace/models/experiments',headers=h['admin']).json()[0]
    assert sum(v['samples'] for v in summary['variants'].values())==2
    assert sum(v['manually_labeled_samples'] for v in summary['variants'].values())==0
    assert c.post(f"/api/workspace/reviews/{first['id']}/moderate",headers=h['admin'],json={'label':first['prediction']['predicted_label'],'reason':'Independent moderation for test case'}).status_code==200
    summary=c.get('/api/workspace/models/experiments',headers=h['admin']).json()[0]
    assert sum(v['manually_labeled_samples'] for v in summary['variants'].values())==1
