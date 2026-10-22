"""Application workflows shared by the web dashboard and mobile app."""
import hashlib
import io
import json
import secrets
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Literal
from fastapi import APIRouter, Depends, HTTPException, Query, BackgroundTasks
from fastapi.responses import StreamingResponse, FileResponse
from pydantic import BaseModel, Field, EmailStr, HttpUrl, model_validator
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError
import models
import extended_models as ex
from database import get_db
from auth import get_current_user, require_roles, hash_password, verify_password
from services.intelligence import account_risk, shop_trust, network_scan, effective_label, latest_prediction
from ml.predictor import predictor, risk_level

router = APIRouter(prefix='/api/workspace', tags=['Nghiệp vụ mở rộng'])
staff = require_roles('admin', 'engineer')
admin = require_roles('admin')


def record(row):
    return {c.name: getattr(row, c.name) for c in row.__table__.columns}


def get_review(db, rid):
    row = db.get(models.Review, rid)
    if not row:
        raise HTTPException(404, 'Không tìm thấy review')
    return row


def own_product(db, pid, user):
    product = db.get(models.Product, pid)
    if not product:
        raise HTTPException(404, 'Không tìm thấy sản phẩm')
    if product.shop_id != user.id and user.role != 'admin':
        raise HTTPException(403, 'Chỉ chủ shop hoặc admin được thực hiện')
    return product


class ProfileInput(BaseModel):
    display_name: str = Field(default='', max_length=150)
    phone: str = Field(default='', max_length=30)
    shop_name: str = Field(default='', max_length=150)
    shop_description: str = Field(default='', max_length=5000)


@router.get('/profile')
def profile(db: Session = Depends(get_db), user=Depends(get_current_user)):
    row = db.get(ex.Profile, user.id)
    return record(row) if row else {'user_id': user.id, 'display_name': '', 'phone': '', 'shop_name': '', 'shop_description': ''}


@router.put('/profile')
def update_profile(body: ProfileInput, db: Session = Depends(get_db), user=Depends(get_current_user)):
    row = db.get(ex.Profile, user.id) or ex.Profile(user_id=user.id)
    values = body.model_dump()
    if user.role not in ('seller', 'admin'):
        values.pop('shop_name'); values.pop('shop_description')
    for key, val in values.items():
        setattr(row, key, val.strip())
    db.add(row); db.commit(); db.refresh(row)
    return record(row)


class PasswordInput(BaseModel):
    current_password: str = Field(max_length=128)
    new_password: str = Field(min_length=8, max_length=128)


@router.post('/password')
def password(body: PasswordInput, db: Session = Depends(get_db), user=Depends(get_current_user)):
    if not verify_password(body.current_password, user.password_hash):
        raise HTTPException(400, 'Mật khẩu hiện tại không đúng')
    user.password_hash = hash_password(body.new_password)
    row = db.get(ex.Profile, user.id) or ex.Profile(user_id=user.id, token_version=0)
    row.token_version = (row.token_version or 0) + 1
    db.add(row); db.commit()
    return {'message': 'Đã đổi mật khẩu. Vui lòng đăng nhập lại.'}


class ForgotInput(BaseModel):
    email: EmailStr


@router.post('/forgot-password')
def forgot(body: ForgotInput, db: Session = Depends(get_db)):
    user = db.query(models.User).filter_by(email=body.email.lower()).first()
    if user and user.is_active:
        token = secrets.token_urlsafe(32)
        db.add(ex.PasswordReset(user_id=user.id, token_hash=hashlib.sha256(token.encode()).hexdigest(),
                                expires_at=datetime.now(timezone.utc)+timedelta(minutes=15)))
        db.add(ex.Outbox(channel='email', recipient=user.email, payload=json.dumps({'kind': 'password_reset', 'token': token})))
        db.commit()
    return {'message': 'Nếu email hợp lệ, hướng dẫn khôi phục đã được đưa vào hàng đợi gửi.'}


class ResetInput(BaseModel):
    token: str = Field(min_length=20, max_length=200)
    new_password: str = Field(min_length=8, max_length=128)


@router.post('/reset-password')
def reset(body: ResetInput, db: Session = Depends(get_db)):
    token_hash = hashlib.sha256(body.token.encode()).hexdigest()
    row = db.query(ex.PasswordReset).filter_by(token_hash=token_hash, used=False).first()
    if not row or row.expires_at.replace(tzinfo=timezone.utc) <= datetime.now(timezone.utc):
        raise HTTPException(400, 'Token không hợp lệ hoặc đã hết hạn')
    user = db.get(models.User, row.user_id)
    row.used = True
    user.password_hash = hash_password(body.new_password)
    profile = db.get(ex.Profile, user.id) or ex.Profile(user_id=user.id, token_version=0)
    profile.token_version = (profile.token_version or 0)+1
    db.add(profile); db.commit()
    return {'message': 'Đã đặt lại mật khẩu'}


class ExtraInput(BaseModel):
    image_url: HttpUrl | None = None
    stock: int = Field(default=0, ge=0)
    sale_status: Literal['active', 'paused', 'sold_out'] = 'active'
    canonical_key: str = Field(default='', max_length=200)


@router.get('/products/{pid}/extra')
def product_extra(pid: int, db: Session = Depends(get_db)):
    row = db.get(ex.ProductExtra, pid)
    return record(row) if row else {'product_id': pid, 'image_url': '', 'stock': 0, 'sale_status': 'active', 'canonical_key': ''}


@router.put('/products/{pid}/extra')
def put_extra(pid: int, body: ExtraInput, db: Session = Depends(get_db), user=Depends(get_current_user)):
    own_product(db, pid, user)
    row = db.get(ex.ProductExtra, pid) or ex.ProductExtra(product_id=pid)
    for key, val in body.model_dump().items():
        setattr(row, key, str(val) if key == 'image_url' and val else '' if key == 'image_url' else val)
    db.add(row); db.commit(); db.refresh(row)
    return record(row)


class CategoryInput(BaseModel):
    name: str = Field(min_length=2, max_length=100)


@router.get('/categories')
def categories(db: Session = Depends(get_db)):
    return [record(x) for x in db.query(ex.Category).order_by(ex.Category.name).all()]


@router.post('/categories')
def add_category(body: CategoryInput, db: Session = Depends(get_db), user=Depends(admin)):
    if db.query(ex.Category).filter_by(name=body.name.strip()).first():
        raise HTTPException(409, 'Danh mục đã tồn tại')
    row = ex.Category(name=body.name.strip()); db.add(row); db.commit(); db.refresh(row)
    return record(row)


@router.put('/categories/{cid}')
def update_category(cid: int, body: CategoryInput, db: Session = Depends(get_db), user=Depends(admin)):
    row = db.get(ex.Category, cid)
    if not row:
        raise HTTPException(404, 'Danh mục không tồn tại')
    existing = db.query(ex.Category).filter_by(name=body.name.strip()).first()
    if existing and existing.id != cid:
        raise HTTPException(409, 'Danh mục đã tồn tại')
    db.query(models.Product).filter_by(category=row.name).update({'category': body.name.strip()})
    row.name = body.name.strip(); db.commit()
    return record(row)


@router.delete('/categories/{cid}')
def delete_category(cid: int, db: Session = Depends(get_db), user=Depends(admin)):
    row = db.get(ex.Category, cid)
    if not row:
        raise HTTPException(404, 'Danh mục không tồn tại')
    if db.query(models.Product).filter_by(category=row.name).first():
        raise HTTPException(409, 'Danh mục đang có sản phẩm')
    db.delete(row); db.commit()
    return {'message': 'Đã xóa danh mục'}


class PurchaseInput(BaseModel):
    product_id: int = Field(gt=0)
    quantity: int = Field(default=1, ge=1, le=100)


@router.post('/purchases')
def purchase(body: PurchaseInput, db: Session = Depends(get_db), user=Depends(get_current_user)):
    product = db.get(models.Product, body.product_id)
    if not product or product.shop_id == user.id:
        raise HTTPException(400, 'Sản phẩm không tồn tại hoặc thuộc shop của bạn')
    extra = db.get(ex.ProductExtra, product.id)
    if extra and (extra.sale_status != 'active' or extra.stock < body.quantity):
        raise HTTPException(409, 'Sản phẩm tạm ngừng bán hoặc không đủ tồn kho')
    row = ex.Purchase(user_id=user.id, product_id=product.id, quantity=body.quantity)
    db.add(row); db.commit(); db.refresh(row)
    return record(row)


@router.get('/purchases')
def purchases(db: Session = Depends(get_db), user=Depends(get_current_user)):
    query = db.query(ex.Purchase)
    if user.role == 'seller':
        query = query.join(models.Product).filter(models.Product.shop_id == user.id)
    elif user.role != 'admin':
        query = query.filter(ex.Purchase.user_id == user.id)
    return [record(x) for x in query.order_by(ex.Purchase.id.desc()).limit(200).all()]


class PurchaseStatus(BaseModel):
    status: Literal['completed', 'cancelled']


@router.patch('/purchases/{oid}')
def complete_purchase(oid: int, body: PurchaseStatus, db: Session = Depends(get_db), user=Depends(get_current_user)):
    row = db.get(ex.Purchase, oid)
    if not row:
        raise HTTPException(404, 'Đơn hàng không tồn tại')
    own_product(db, row.product_id, user)
    if row.status != 'pending':
        raise HTTPException(409, 'Đơn đã được xử lý')
    transitioned=db.query(ex.Purchase).filter_by(id=oid,status='pending').update({'status':body.status})
    if transitioned!=1:
        raise HTTPException(409,'Đơn vừa được xử lý bởi một yêu cầu khác')
    extra = db.get(ex.ProductExtra, row.product_id)
    if body.status == 'completed' and extra:
        changed = db.query(ex.ProductExtra).filter(ex.ProductExtra.product_id==row.product_id, ex.ProductExtra.stock>=row.quantity).update({ex.ProductExtra.stock: ex.ProductExtra.stock-row.quantity})
        if changed != 1:
            db.rollback()
            raise HTTPException(409, 'Không đủ tồn kho')
    db.commit(); db.refresh(row)
    return record(row)


class TextInput(BaseModel):
    content: str = Field(min_length=1, max_length=5000)
    model_version: str | None = Field(default=None,max_length=100)


@router.post('/analyze')
def analyze(body: TextInput, db: Session = Depends(get_db), user=Depends(get_current_user)):
    if not body.content.strip():
        raise HTTPException(422, 'Nội dung không được rỗng')
    result = predictor.predict(body.content)
    db.add(ex.Analysis(user_id=user.id, content=body.content.strip(), result=json.dumps(result)))
    db.commit()
    return {**result, 'risk_level': risk_level(result['fake_probability'])}


@router.get('/analyses')
def analyses(db: Session = Depends(get_db), user=Depends(get_current_user)):
    return [{**record(row), 'result': json.loads(row.result)} for row in db.query(ex.Analysis).filter_by(user_id=user.id).order_by(ex.Analysis.id.desc()).limit(200).all()]


@router.post('/explain')
def explain(body: TextInput, user=Depends(get_current_user)):
    if not body.content.strip():
        raise HTTPException(422, 'Nội dung không được rỗng')
    selected=predictor
    if body.model_version and body.model_version!=predictor.metadata.get('version'):
        from ml.predictor import version_predictor
        selected=None
        for path in (Path(__file__).resolve().parents[1]/'ml/models').glob('*.json'):
            metadata=json.loads(path.read_text(encoding='utf-8'))
            if isinstance(metadata,dict) and metadata.get('version')==body.model_version and path.with_suffix('.joblib').exists():
                selected=version_predictor(path.with_suffix('.joblib').name);break
        if selected is None:raise HTTPException(409,'Không còn artifact của phiên bản đã phân tích review này.')
    return selected.explain(body.content)


@router.get('/reviews/{rid}/discussion')
def discussion(rid: int, db: Session = Depends(get_db)):
    get_review(db, rid)
    return {'replies': [record(x) for x in db.query(ex.ReviewReply).filter_by(review_id=rid).all()],
        'useful': db.query(ex.ReviewVote).filter_by(review_id=rid, useful=True).count(),
        'not_useful': db.query(ex.ReviewVote).filter_by(review_id=rid, useful=False).count(),
        'effective_label': effective_label(db, rid),
        'moderation': record(db.query(ex.Moderation).filter_by(review_id=rid).order_by(ex.Moderation.id.desc()).first()) if db.query(ex.Moderation).filter_by(review_id=rid).first() else None}


@router.post('/reviews/{rid}/reply')
def reply(rid: int, body: TextInput, db: Session = Depends(get_db), user=Depends(get_current_user)):
    review = get_review(db, rid); own_product(db, review.product_id, user)
    row = ex.ReviewReply(review_id=rid, user_id=user.id, content=body.content.strip())
    if not row.content:
        raise HTTPException(422, 'Nội dung không được rỗng')
    db.add(row); db.commit(); db.refresh(row)
    return record(row)


class VoteInput(BaseModel):
    useful: bool


@router.put('/reviews/{rid}/vote')
def vote(rid: int, body: VoteInput, db: Session = Depends(get_db), user=Depends(get_current_user)):
    review = get_review(db,rid)
    if review.user_id == user.id:
        raise HTTPException(400, 'Không tự bình chọn review của mình')
    row = db.query(ex.ReviewVote).filter_by(review_id=rid,user_id=user.id).first() or ex.ReviewVote(review_id=rid,user_id=user.id)
    row.useful=body.useful; db.add(row); db.commit()
    return discussion(rid, db)


class ModerateInput(BaseModel):
    label: Literal['fake', 'real']
    reason: str = Field(min_length=5, max_length=2000)


@router.post('/reviews/{rid}/moderate')
def moderate(rid: int, body: ModerateInput, db: Session = Depends(get_db), user=Depends(admin)):
    review = get_review(db,rid)
    db.add(ex.Moderation(review_id=rid,admin_id=user.id,label=body.label,reason=body.reason))
    db.add(ex.Notification(user_id=review.user_id,message=f'Review #{rid} được kiểm duyệt: {body.label}. {body.reason}'))
    db.commit()
    return {'message':'Đã lưu nhãn kiểm duyệt; giữ nguyên kết quả AI để đối chiếu.'}


@router.get('/reviews/{rid}/audit')
def review_audit(rid: int, db: Session = Depends(get_db), user=Depends(staff)):
    get_review(db, rid)
    return {'predictions': [record(x) for x in db.query(models.Prediction).filter_by(review_id=rid).order_by(models.Prediction.id.desc()).all()],
            'moderations': [record(x) for x in db.query(ex.Moderation).filter_by(review_id=rid).order_by(ex.Moderation.id.desc()).all()]}


@router.get('/reports')
def reports(db: Session = Depends(get_db), user=Depends(get_current_user)):
    query=db.query(models.ReviewReport)
    if user.role not in ('admin','engineer'):
        query=query.filter_by(reporter_user_id=user.id)
    return [{**record(x), 'resolution': record(db.get(ex.ReportResolution,x.id)) if db.get(ex.ReportResolution,x.id) else None} for x in query.order_by(models.ReviewReport.id.desc()).limit(200).all()]


class ResolveInput(BaseModel):
    status: Literal['investigating','resolved','rejected']
    note: str = Field(min_length=5,max_length=2000)


@router.patch('/reports/{rid}')
def resolve_report(rid:int,body:ResolveInput,db:Session=Depends(get_db),user=Depends(admin)):
    report=db.get(models.ReviewReport,rid)
    if not report:
        raise HTTPException(404,'Báo cáo không tồn tại')
    report.status=body.status
    row=db.get(ex.ReportResolution,rid) or ex.ReportResolution(report_id=rid)
    row.admin_id=user.id; row.note=body.note
    db.add(row); db.add(ex.Notification(user_id=report.reporter_user_id,message=f'Báo cáo #{rid}: {body.status}. {body.note}'))
    db.commit()
    return record(report)


@router.get('/accounts/{uid}/risk')
def risk(uid:int,db:Session=Depends(get_db),user=Depends(get_current_user)):
    if user.id != uid and user.role not in ('admin','engineer'):
        raise HTTPException(403,'Không có quyền xem phân tích tài khoản khác')
    return account_risk(db,uid)


@router.get('/accounts/risk/all')
def all_risks(db:Session=Depends(get_db),user=Depends(staff)):
    return [account_risk(db,u.id) for u in db.query(models.User).all()]


@router.post('/network/scan')
def scan(db:Session=Depends(get_db),user=Depends(staff)):
    for uid, in db.query(models.Review.user_id).distinct().all():
        account_risk(db,uid,persist=True)
    return network_scan(db,persist=True)


@router.get('/network')
def network(db:Session=Depends(get_db),user=Depends(staff)):
    return network_scan(db)


@router.get('/shops')
def shops(db:Session=Depends(get_db)):
    return [shop_trust(db,u.id) for u in db.query(models.User).filter(models.User.role.in_(['seller','admin'])).all() if db.query(models.Product.id).filter_by(shop_id=u.id).first()]


@router.get('/shops/{sid}/trust')
def trust(sid:int,db:Session=Depends(get_db)):
    return shop_trust(db,sid)


@router.post('/shops/{sid}/snapshot')
def snapshot(sid:int,db:Session=Depends(get_db),user=Depends(get_current_user)):
    if user.id!=sid and user.role not in ('admin','engineer'):
        raise HTTPException(403,'Không có quyền chụp thống kê shop này')
    return shop_trust(db,sid,True)


@router.get('/compare')
def compare(canonical_key:str=Query(min_length=1,max_length=200),db:Session=Depends(get_db)):
    rows=db.query(models.Product).join(ex.ProductExtra).filter(ex.ProductExtra.canonical_key==canonical_key).all()
    return [{'product_id':p.id,'name':p.name,'price':p.price,'shop':shop_trust(db,p.shop_id)} for p in rows]


class ThresholdInput(BaseModel):
    normal_max: float = Field(ge=0,le=100)
    suspicious_min: float = Field(ge=0,le=100)
    @model_validator(mode='after')
    def ordered(self):
        if self.normal_max>=self.suspicious_min:
            raise ValueError('Ngưỡng bình thường phải nhỏ hơn ngưỡng đáng ngờ')
        return self


@router.get('/settings')
def settings_view(db:Session=Depends(get_db),user=Depends(staff)):
    from services.intelligence import levels
    return {'risk_thresholds':levels(db),'blacklist':json.loads(db.get(ex.Setting,'blacklist').value) if db.get(ex.Setting,'blacklist') else []}


@router.put('/settings/thresholds')
def set_thresholds(body:ThresholdInput,db:Session=Depends(get_db),user=Depends(admin)):
    row=db.get(ex.Setting,'risk_thresholds') or ex.Setting(key='risk_thresholds')
    row.value=body.model_dump_json(); db.add(row); db.commit()
    return body


class BlacklistInput(BaseModel):
    keywords: list[str] = Field(max_length=200)


@router.put('/settings/blacklist')
def set_blacklist(body:BlacklistInput,db:Session=Depends(get_db),user=Depends(admin)):
    if any(not k.strip() or len(k)>100 for k in body.keywords):
        raise HTTPException(422,'Từ khóa phải có 1–100 ký tự')
    row=db.get(ex.Setting,'blacklist') or ex.Setting(key='blacklist')
    row.value=json.dumps(body.keywords,ensure_ascii=False);db.add(row);db.commit()
    return body


@router.put('/products/{pid}/watch')
def watch(pid:int,body:VoteInput,db:Session=Depends(get_db),user=Depends(get_current_user)):
    if not db.get(models.Product,pid):
        raise HTTPException(404,'Sản phẩm không tồn tại')
    row=db.query(ex.Watch).filter_by(product_id=pid,user_id=user.id).first()
    if body.useful and not row:
        db.add(ex.Watch(product_id=pid,user_id=user.id))
    elif not body.useful and row:
        db.delete(row)
    db.commit();return {'watching':body.useful}


@router.get('/notifications')
def notifications(db:Session=Depends(get_db),user=Depends(get_current_user)):
    return [record(x) for x in db.query(ex.Notification).filter_by(user_id=user.id).order_by(ex.Notification.id.desc()).limit(100).all()]


@router.patch('/notifications/{nid}/read')
def notification_read(nid:int,db:Session=Depends(get_db),user=Depends(get_current_user)):
    row=db.get(ex.Notification,nid)
    if not row or row.user_id!=user.id:
        raise HTTPException(404,'Thông báo không tồn tại')
    row.read=True;db.commit();return {'message':'Đã đọc'}


@router.get('/models')
def model_info(user=Depends(staff)):
    root=Path(__file__).resolve().parents[1]/'ml/models'
    candidates=[]
    for file in root.glob('*.json'):
        try:
            metadata=json.loads(file.read_text(encoding='utf-8'))
            if isinstance(metadata,dict) and 'model_name' in metadata:
                candidates.append({'file':file.name,'metrics':metadata})
        except (ValueError,OSError):
            pass
    return {'ready':predictor.ready,'active':predictor.metadata,'error':predictor.error,'experiments':candidates}


@router.get('/models/figure/{filename}')
def model_figure(filename:Literal['distilbert_curves.png','calibration.png']):
    path=Path(__file__).resolve().parents[1]/'ml/models'/filename
    if not path.exists():raise HTTPException(404,'Chưa có hình thực nghiệm này')
    return FileResponse(path,media_type='image/png')


@router.get('/content')
def content(kind: str|None=None,db:Session=Depends(get_db)):
    q=db.query(ex.Content)
    if kind:q=q.filter_by(kind=kind)
    return [record(x) for x in q.all()]


class ContentInput(BaseModel):
    kind:Literal['faq','banner','guide','policy']
    title:str=Field(min_length=2,max_length=200)
    body:str=Field(min_length=2,max_length=10000)


@router.post('/content')
def add_content(body:ContentInput,db:Session=Depends(get_db),user=Depends(admin)):
    row=ex.Content(**body.model_dump());db.add(row);db.commit();db.refresh(row);return record(row)


@router.put('/content/{cid}')
def update_content(cid:int,body:ContentInput,db:Session=Depends(get_db),user=Depends(admin)):
    row=db.get(ex.Content,cid)
    if not row:raise HTTPException(404,'Nội dung không tồn tại')
    for k,v in body.model_dump().items():setattr(row,k,v)
    db.commit();return record(row)


@router.delete('/content/{cid}')
def delete_content(cid:int,db:Session=Depends(get_db),user=Depends(admin)):
    row=db.get(ex.Content,cid)
    if not row:raise HTTPException(404,'Nội dung không tồn tại')
    db.delete(row);db.commit();return {'message':'Đã xóa'}


@router.get('/search/health')
def search_health(user=Depends(staff)):
    from services.search import health
    return health()


@router.post('/search/reindex')
def search_reindex(db:Session=Depends(get_db),user=Depends(staff)):
    from services.search import reindex
    return reindex(db)


@router.get('/search')
def search(q:str=Query('',max_length=200),kind:Literal['product','review']|None=None,label:Literal['fake','real']|None=None,
           page:int=Query(1,ge=1),page_size:int=Query(20,ge=1,le=100),user=Depends(get_current_user)):
    from services.search import search as do_search
    return do_search(q,kind,label,page,page_size)


def training_worker(job_id):
    import os
    import subprocess
    import sys
    from database import SessionLocal
    root=Path(__file__).resolve().parents[1]
    with SessionLocal() as db:
        row=db.get(ex.TrainingJob,job_id)
        row.status='running';db.commit()
        output=root/'ml/models'/f'candidate_{job_id}.joblib'
        log=root/'ml/models'/f'training_{job_id}.log'
        try:
            arguments=[]
            if row.kind=='svm_feedback':
                from ml.feedback import export_feedback
                feedback_path=root/'ml/models'/f'feedback_{job_id}.csv'
                if export_feedback(db,feedback_path)==0:
                    raise ValueError('Chưa có nhãn kiểm duyệt để huấn luyện từ phản hồi')
                arguments=['--feedback',str(feedback_path)]
            with log.open('w',encoding='utf-8') as stream:
                result=subprocess.run([sys.executable,'-m','ml.train_bert' if row.kind=='distilbert' else 'ml.train','--output',str(output),*arguments],cwd=root,
                                      stdout=stream,stderr=subprocess.STDOUT,timeout=6*3600,env={**os.environ,'PYTHONIOENCODING':'utf-8'})
            row.status='completed' if result.returncode==0 else 'failed'
            row.details=json.dumps({'artifact':output.name,'log':log.name,'returncode':result.returncode})
        except Exception as exc:
            row.status='failed';row.details=str(exc)
        db.commit()


class TrainingInput(BaseModel):
    kind:Literal['svm','svm_feedback','distilbert']='svm'


@router.post('/models/train')
def train_model(body:TrainingInput,tasks:BackgroundTasks,db:Session=Depends(get_db),user=Depends(staff)):
    if db.query(ex.TrainingJob).filter(ex.TrainingJob.status.in_(['queued','running'])).first():
        raise HTTPException(409,'Đang có tác vụ huấn luyện. Chờ hoàn tất trước khi tạo tác vụ mới.')
    row=ex.TrainingJob(requested_by=user.id,kind=body.kind);db.add(row);db.commit();db.refresh(row)
    tasks.add_task(training_worker,row.id)
    return record(row)


@router.get('/models/jobs')
def jobs(db:Session=Depends(get_db),user=Depends(staff)):
    return [record(x) for x in db.query(ex.TrainingJob).order_by(ex.TrainingJob.id.desc()).limit(50).all()]


@router.get('/models/jobs/{jid}/log')
def training_log(jid:int,db:Session=Depends(get_db),user=Depends(staff)):
    row=db.get(ex.TrainingJob,jid)
    if not row:raise HTTPException(404,'Không có tác vụ')
    path=Path(__file__).resolve().parents[1]/'ml/models'/f'training_{jid}.log'
    return {'log':path.read_text(encoding='utf-8',errors='replace')[-20000:] if path.exists() else ''}


@router.post('/models/jobs/{jid}/activate')
def activate(jid:int,db:Session=Depends(get_db),user=Depends(admin)):
    from config import settings
    job=db.get(ex.TrainingJob,jid)
    if not job or job.status!='completed':raise HTTPException(409,'Tác vụ chưa hoàn tất thành công')
    path=Path(__file__).resolve().parents[1]/'ml/models'/f'candidate_{jid}.joblib'
    previous=settings.model_path;settings.model_path=str(path)
    if not predictor.reload():
        settings.model_path=previous;raise HTTPException(400,'Model mới không hợp lệ; tiếp tục giữ model hiện tại')
    row=db.get(ex.Setting,'active_model') or ex.Setting(key='active_model')
    row.value=str(path);db.add(row);db.commit()
    return {'active':predictor.metadata}


@router.get('/exports/{fmt}')
def export(fmt:Literal['xlsx','pdf'],kind:Literal['reviews','shops','categories','clusters','models']='reviews',db:Session=Depends(get_db),user=Depends(require_roles('seller','admin','engineer'))):
    from services.reporting import report_rows
    rows=report_rows(db,user,kind)
    buffer=io.BytesIO()
    if fmt=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb=Workbook();ws=wb.active;ws.title=kind.title()
        for row in rows:
            ws.append(row)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(wrap_text=True,vertical='top')
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1E3A5F')
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for col in ['A','B','C','D','E']:ws.column_dimensions[col].width=18
        ws.column_dimensions['F'].width=90
        wb.save(buffer);mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        from reportlab.platypus import SimpleDocTemplate,Paragraph,LongTable,TableStyle,Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4,landscape
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from xml.sax.saxutils import escape
        font=Path('C:/Windows/Fonts/arial.ttf')
        if font.exists():pdfmetrics.registerFont(TTFont('ReportFont',str(font)))
        styles=getSampleStyleSheet();style=styles['BodyText'];style.fontName='ReportFont' if font.exists() else 'Helvetica';style.fontSize=8
        widths=[48,52,48,40,72,420] if kind=='reviews' else [780/len(rows[0])]*len(rows[0])
        table=LongTable([[Paragraph(escape(str(v)),style) for v in row] for row in rows],repeatRows=1,splitInRow=1,colWidths=widths)
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DCE6F1')),('GRID',(0,0),(-1,-1),.3,colors.grey),('VALIGN',(0,0),(-1,-1),'TOP')]))
        SimpleDocTemplate(buffer,pagesize=landscape(A4),leftMargin=24,rightMargin=24).build([Paragraph('Fake Review Detection — Review report',styles['Title']),Spacer(1,15),table]);mime='application/pdf'
    buffer.seek(0)
    return StreamingResponse(buffer,media_type=mime,headers={'Content-Disposition':f'attachment; filename="{kind}-report.{fmt}"'})


class ActivateInput(BaseModel):
    filename:str=Field(min_length=1,max_length=200)


@router.post('/models/activate')
def activate_version(body:ActivateInput,db:Session=Depends(get_db),user=Depends(admin)):
    from ml.predictor import version_predictor
    from config import settings
    candidate=version_predictor(body.filename)
    predictor.model,predictor.metadata,predictor.error=candidate.model,candidate.metadata,None
    settings.model_path=str(candidate.path)
    row=db.get(ex.Setting,'active_model') or ex.Setting(key='active_model')
    row.value='ml/models/'+body.filename;db.add(row);db.commit()
    return {'active':predictor.metadata}


class ExperimentInput(BaseModel):
    model_a:str=Field(max_length=200)
    model_b:str=Field(max_length=200)
    traffic_b:int=Field(ge=1,le=99)


@router.post('/models/experiments')
def start_experiment(body:ExperimentInput,db:Session=Depends(get_db),user=Depends(admin)):
    from ml.predictor import version_predictor
    if body.model_a==body.model_b:raise HTTPException(422,'Hai phiên bản phải khác nhau')
    version_predictor(body.model_a);version_predictor(body.model_b)
    db.query(ex.ModelExperiment).filter_by(active=True).update({'active':False})
    row=ex.ModelExperiment(**body.model_dump());db.add(row);db.commit();db.refresh(row)
    return record(row)


@router.get('/models/experiments')
def experiments(db:Session=Depends(get_db),user=Depends(staff)):
    output=[]
    for row in db.query(ex.ModelExperiment).order_by(ex.ModelExperiment.id.desc()).limit(20):
        observations=db.query(ex.ModelObservation).filter_by(experiment_id=row.id).all()
        variants={}
        for variant in ('A','B'):
            rows=[x for x in observations if x.variant==variant]
            judged=[]
            for item in rows:
                manual=db.query(ex.Moderation).filter_by(review_id=item.review_id).order_by(ex.Moderation.id.desc()).first()
                if manual:judged.append(json.loads(item.result)['predicted_label']==manual.label)
            variants[variant]={'samples':len(rows),'mean_latency_ms':sum(x.latency_ms for x in rows)/len(rows) if rows else None,
                'manually_labeled_samples':len(judged),'agreement_with_moderator':sum(judged)/len(judged) if judged else None}
        output.append({**record(row),'variants':variants})
    return output


@router.post('/models/experiments/{eid}/stop')
def stop_experiment(eid:int,db:Session=Depends(get_db),user=Depends(admin)):
    row=db.get(ex.ModelExperiment,eid)
    if not row:raise HTTPException(404,'Không có thí nghiệm')
    row.active=False;db.commit();return record(row)


@router.get('/shops/{sid}/advice')
def advice(sid:int,db:Session=Depends(get_db),user=Depends(get_current_user)):
    if user.id!=sid and user.role not in ('admin','engineer'):
        raise HTTPException(403,'Không có quyền xem phản hồi nội bộ shop này')
    reviews=db.query(models.Review).join(models.Product).filter(models.Product.shop_id==sid).all()
    genuine=[r for r in reviews if effective_label(db,r.id)=='real']
    negative=[r for r in genuine if r.rating<=3]
    suggestions=[]
    if negative:suggestions.append({'action':'Ưu tiên phản hồi khách hàng có trải nghiệm chưa tốt','review_ids':[r.id for r in negative]})
    pending_reports=db.query(models.ReviewReport).join(models.Review).join(models.Product).filter(models.Product.shop_id==sid,models.ReviewReport.status.in_(['pending','investigating'])).all()
    if pending_reports:suggestions.append({'action':'Đối chiếu các review đang có báo cáo với quản trị viên','report_ids':[r.id for r in pending_reports]})
    if not genuine:suggestions.append({'action':'Chưa đủ review được mô hình/kiểm duyệt xem là gốc để tổng hợp trải nghiệm.'})
    return {'shop_id':sid,'genuine_reviews':len(genuine),'low_rating_genuine_reviews':len(negative),'suggestions':suggestions,'method':'Quy tắc dựa trên dữ liệu; không phải khuyến nghị do mô hình ngôn ngữ sinh ra.'}
